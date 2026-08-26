import base64
import csv
import io
import logging

from odoo import models, fields, api, _
from odoo.exceptions import UserError

logger = logging.getLogger(__name__)

CSV_DELIMITER = ';'
# intestazioni accettate nella prima riga del file, in italiano e inglese
HEADER_NAMES = ['nome', 'name']
HEADER_CODES = ['codice', 'code', 'tag_code']


class CaImportaBadgeGenerico(models.TransientModel):
    _name = 'ca.importa_badge_generico'
    _description = 'Import Generic Badges'

    state = fields.Selection([
        ('draft', 'Template'),
        ('template', 'Template Ready'),
        ('done', 'Imported')
    ], default='draft', required=True)

    tipo_badge_id = fields.Many2one(
        'ca.tipo_badge_generico', string="Generic Badge Type", required=True)
    quantity = fields.Integer(string="Badges to Create", default=10)
    file_format = fields.Selection([
        ('csv', 'CSV'),
        ('xlsx', 'Excel (xlsx)')
    ], default='csv', required=True, string="Template Format")

    template_file = fields.Binary(string="Template", readonly=True)
    template_filename = fields.Char(readonly=True)

    import_file = fields.Binary(string="Filled File")
    import_filename = fields.Char()

    result_message = fields.Text(readonly=True)
    created_tag_ids = fields.Many2many('ca.tag', readonly=True)

    @api.onchange('tipo_badge_id', 'quantity', 'file_format')
    def _onchange_template_data(self):
        """Il template gia' generato non vale piu' se cambiano i dati."""
        for record in self:
            if record.state == 'template':
                record.state = 'draft'
                record.template_file = False
                record.template_filename = False

    def _reload_action(self):
        """Ricarica il wizard mantenendo aperta la finestra."""
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    def _build_csv(self, names):
        """CSV separato da ';' con BOM: Excel apre le colonne correttamente."""
        output = io.StringIO()
        writer = csv.writer(output, delimiter=CSV_DELIMITER,
                            quoting=csv.QUOTE_MINIMAL, lineterminator='\r\n')
        writer.writerow(['Nome', 'Codice'])
        for name in names:
            writer.writerow([name, ''])
        return output.getvalue().encode('utf-8-sig')

    def _build_xlsx(self, names):
        import xlsxwriter
        output = io.BytesIO()
        workbook = xlsxwriter.Workbook(output, {'in_memory': True})
        sheet = workbook.add_worksheet('Badge')
        header = workbook.add_format({'bold': True})
        text = workbook.add_format({'num_format': '@'})
        sheet.write(0, 0, 'Nome', header)
        sheet.write(0, 1, 'Codice', header)
        sheet.set_column(0, 0, 30)
        sheet.set_column(1, 1, 25, text)
        for row, name in enumerate(names, start=1):
            sheet.write_string(row, 0, name)
            sheet.write_blank(row, 1, None, text)
        workbook.close()
        return output.getvalue()

    def action_generate_template(self):
        """Genera il file con la colonna Nome gia' compilata."""
        self.ensure_one()
        if self.quantity <= 0:
            raise UserError(_('The number of badges must be greater than zero'))
        names = self.tipo_badge_id.get_next_badge_names(self.quantity)
        if not names:
            raise UserError(
                _('Cannot build the badge names for this type, '
                  'check the badge name prefix'))
        if self.file_format == 'xlsx':
            content = self._build_xlsx(names)
            extension = 'xlsx'
        else:
            content = self._build_csv(names)
            extension = 'csv'
        prefix = self.tipo_badge_id.get_badge_prefix().replace(' ', '_')
        self.write({
            'state': 'template',
            'template_file': base64.b64encode(content),
            'template_filename': f"badge_{prefix}_{names[0].split(' - ')[-1]}"
                                 f"_{len(names)}.{extension}",
        })
        return self._reload_action()

    def _read_rows(self):
        """Righe (nome, codice) del file caricato, intestazione esclusa."""
        self.ensure_one()
        content = base64.b64decode(self.import_file)
        filename = (self.import_filename or '').lower()
        if filename.endswith('.xlsx') or filename.endswith('.xlsm'):
            rows = self._read_rows_xlsx(content)
        elif filename.endswith('.xls'):
            raise UserError(
                _('The old .xls format is not supported, save the file '
                  'as .xlsx or .csv'))
        else:
            rows = self._read_rows_csv(content)
        if rows and len(rows[0]) > 1:
            first_name = (rows[0][0] or '').strip().lower()
            first_code = (rows[0][1] or '').strip().lower()
            if first_name in HEADER_NAMES and first_code in HEADER_CODES:
                rows = rows[1:]
        return rows

    def _read_rows_csv(self, content):
        try:
            text = content.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = content.decode('latin-1')
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=';,\t')
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = CSV_DELIMITER
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        return [[cell for cell in row] for row in reader]

    def _read_rows_xlsx(self, content):
        from openpyxl import load_workbook
        workbook = load_workbook(
            filename=io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = []
        for row in sheet.iter_rows(values_only=True):
            rows.append([
                '' if cell is None else str(cell).strip() for cell in row])
        workbook.close()
        return rows

    def _parse_rows(self, rows):
        """Valida tutte le righe: o si importa tutto o non si importa nulla.

        Un import parziale lascerebbe badge gia' scritti sui lettori
        fisici, difficili da ripulire a mano.
        """
        tag_model = self.env['ca.tag'].with_context(active_test=False)
        tipo_model = self.env['ca.tipo_badge_generico']
        badge_vals = []
        errors = []
        skipped = 0
        seen_codes = {}
        seen_names = {}
        for index, row in enumerate(rows, start=1):
            name = (row[0] if len(row) > 0 else '') or ''
            code = (row[1] if len(row) > 1 else '') or ''
            name = str(name).strip()
            code = tipo_model.normalize_tag_code(str(code))
            if not name and not code:
                continue
            if not code:
                skipped += 1
                continue
            if not name:
                errors.append(
                    _('Row %s: the code %s has no badge name') % (index, code))
                continue
            if code in seen_codes:
                errors.append(
                    _('Row %s: the code %s is already used on row %s of '
                      'the file') % (index, code, seen_codes[code]))
                continue
            seen_codes[code] = index
            if not tipo_model.is_valid_tag_code(code):
                errors.append(
                    _('Row %s: the code %s is not 16 hexadecimal characters, '
                      'format the code column as text and write it again') % (
                        index, code))
                continue
            if name in seen_names:
                errors.append(
                    _('Row %s: the badge name %s is already used on row %s '
                      'of the file') % (index, name, seen_names[name]))
                continue
            seen_names[name] = index
            existing = tag_model.search([('tag_code', '=', code)], limit=1)
            if existing:
                errors.append(
                    _('Row %s: the code %s already exists on badge %s') % (
                        index, code, existing.name))
                continue
            existing = tag_model.search([('name', '=', name)], limit=1)
            if existing:
                errors.append(
                    _('Row %s: the badge name %s already exists with '
                      'code %s') % (index, name, existing.tag_code))
                continue
            badge_vals.append({'name': name, 'tag_code': code})
        if errors:
            raise UserError(
                _('The file cannot be imported, no badge has been created:')
                + '\n' + '\n'.join(errors))
        if not badge_vals:
            raise UserError(_('No badge code found in the file'))
        return badge_vals, skipped

    def action_import(self):
        """Crea in blocco i badge del file e li scrive sui lettori di sede."""
        self.ensure_one()
        if not self.import_file:
            raise UserError(_('Load the filled file first'))
        badge_vals, skipped = self._parse_rows(self._read_rows())
        tags = self.tipo_badge_id.create_badges(badge_vals)
        message = _('%s generic badges created') % len(tags)
        if skipped:
            message += '\n' + _('%s rows skipped, code column empty') % skipped
        self.write({
            'state': 'done',
            'result_message': message,
            'created_tag_ids': [(6, 0, tags.ids)],
        })
        logger.info(f"Generic badges import: {message}")
        return self._reload_action()
